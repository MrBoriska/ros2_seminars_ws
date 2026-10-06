import docx

doc = docx.Document('/tmp/file_attachments/илья зыбин, модуль 2/[ROS2-2026] Модуль 2 - Планирование - Вариант №10 - Зыбин Илья.docx')

print("--- Paragraphs ---")
for para in doc.paragraphs:
    print(para.text)

print("--- Tables ---")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                print(para.text)
